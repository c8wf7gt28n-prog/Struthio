#!/usr/bin/env python3
"""STRUTHIO HANDHELD · Build Manual 1.1 -> 1.2 (beginner expansion).

Takes the owner's STRUTHIO11.docx and inserts new sections in its own styles
(Heading 1/2, Audio Guide Title/Body, Source Note, shaded tables, callouts,
centred figures with grey italic captions). Nothing of 1.1 is removed; only the
title-page edition number changes.
    python3 tools/manual/expand_v12.py IN.docx FIG_DIR OUT.docx
FIG_DIR holds the figures from figs12.py plus service.png (service-mode screen
rendered on the host from firmware/main/service.c's layout).
"""
import copy, os, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, FIG, OUT = sys.argv[1:4]
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..'))
def R(*p): return os.path.join(ROOT, *p)
def FG(n): return os.path.join(FIG, n)

doc = Document(SRC)
body = doc.element.body

def el_text(e):
    return ''.join(t.text or '' for t in e.iter(qn('w:t')))

def find_heading(text, style=None):
    for e in body.iterchildren():
        if e.tag == qn('w:p') and el_text(e).strip() == text:
            return e
    raise SystemExit('anchor not found: ' + text)

def before_break(h):
    """the page-break paragraph that opens a chapter, if there is one"""
    prev = h.getprevious()
    if prev is not None and prev.tag == qn('w:p') and prev.find('.//' + qn('w:br')) is not None and not el_text(prev).strip():
        return prev
    return h

def shd(el_pr, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    el_pr.append(s)

def borders(ppr, color):
    b = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        x = OxmlElement('w:' + side); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '6'); x.set(qn('w:color'), color)
        b.append(x)
    ppr.append(b)

class At:
    """inserts new blocks before an anchor element, in order"""
    def __init__(self, anchor): self.anchor = anchor
    def _put(self, el): self.anchor.addprevious(el); return el
    def _para(self, style=None):
        p = doc.add_paragraph(style=style)
        self._put(p._p)
        return p
    def pagebreak(self):
        p = self._para(); p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
    def h1(self, t, newpage=True):
        if newpage: self.pagebreak()
        self._para('Heading 1').add_run(t)
    def h2(self, t): self._para('Heading 2').add_run(t)
    def p(self, t, italic=False, bold_lead=None):
        p = self._para()
        if bold_lead:
            r = p.add_run(bold_lead); r.bold = True
        r = p.add_run(t); r.italic = italic
        return p
    def bullets(self, items, mark='\u2022 '):
        for it in items:
            p = self._para()
            ppr = p._p.get_or_add_pPr()
            sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '40'); ppr.append(sp)
            ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '288'); ppr.append(ind)
            if isinstance(it, tuple):
                r = p.add_run(mark + it[0]); r.bold = True; r.font.size = Pt(10.5)
                r = p.add_run(it[1]); r.font.size = Pt(10.5)
            else:
                r = p.add_run(mark + it); r.font.size = Pt(10.5)
    def steps(self, items):
        for i, it in enumerate(items, 1):
            self.bullets([it if isinstance(it, tuple) else it], mark=f'{i}. ')
    def checks(self, items):
        for it in items:
            self._para().add_run('\u2610 ' + it)
    def audio(self, title, paras):
        p = self._para('Audio Guide Title'); ppr = p._p.get_or_add_pPr(); shd(ppr, '17365D'); borders(ppr, '17365D')
        p.add_run(title)
        for t in paras:
            p = self._para('Audio Guide Body'); ppr = p._p.get_or_add_pPr(); shd(ppr, 'EAF2F8'); borders(ppr, 'B4C6E7')
            p.add_run(t)
    def note(self, t, fill='FFF2CC', border='D6B656'):
        p = self._para('Source Note'); ppr = p._p.get_or_add_pPr(); shd(ppr, fill); borders(ppr, border)
        p.add_run(t)
    def source(self, t): self._para('Source Note').add_run(t)
    def _table(self, nrows, widths):
        total = 10224
        ws = [int(total * w / sum(widths)) for w in widths]
        t = doc.add_table(rows=nrows, cols=len(widths))
        tbl = t._tbl
        tblPr = tbl.tblPr
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
        for row in t.rows:
            for c, w in zip(row.cells, ws):
                c.width = w * 635  # dxa -> EMU
        self._put(tbl)
        return t, ws
    def table(self, header, rows, widths=None, fill='D9EAD3', size=9.5, bold_first=False):
        widths = widths or [1] * len(header)
        t, ws = self._table(len(rows) + 1, widths)
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; shd(c._tc.get_or_add_tcPr(), fill)
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                r = t.rows[i].cells[j].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
                if bold_first and j == 0: r.bold = True
        for row in t.rows:   # never split a row across pages
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        # keep the header with the first row; repeat it on a new page
        trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
        self._para()  # spacing after the table, as python-docx tables butt into the next block
        return t
    def callout(self, t, fill='FFF2CC', size=10.5):
        tb, _ = self._table(1, [1])
        c = tb.rows[0].cells[0]; shd(c._tc.get_or_add_tcPr(), fill)
        r = c.paragraphs[0].add_run(t); r.bold = True; r.font.size = Pt(size)
        self._para()
    def code(self, lines, size=8.5):
        tb, _ = self._table(1, [1])
        c = tb.rows[0].cells[0]; shd(c._tc.get_or_add_tcPr(), 'F2F2F2')
        p = c.paragraphs[0]
        for i, ln in enumerate(lines):
            r = p.add_run(ln); r.font.name = 'Consolas'; r.font.size = Pt(size)
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Consolas')
            if i < len(lines) - 1: r.add_break()
        self._para()
    def figure(self, path, caption, width=6.4):
        p = self._para(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(width))
        c = self._para(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(8)
        r = c.add_run(caption); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)

# ------------------------------------------------------------------------------------------------
# Title page: edition number
for e in body.iterchildren():
    if e.tag == qn('w:p') and 'STRUTHIO HANDHELD BUILD MANUAL 1.1' in el_text(e):
        for t in e.iter(qn('w:t')):
            if '1.1' in (t.text or ''): t.text = t.text.replace('MANUAL 1.1', 'MANUAL 1.2')
        break

# ------------------------------------------------------------------------------------------------
# 1. "What this 1.2 edition adds" + reading map, before Step 1
a = At(before_break(find_heading('Step 1 \u2014 Understand what you are building')))
a.h1('What this 1.2 edition adds')
a.p('Edition 1.1 told you what to do and in which order. Edition 1.2 keeps every page of 1.1 and adds the "how" '
    'and the "why" for someone who has never soldered, never used a terminal, and never flashed a microcontroller. '
    'All new numbers come from the project files themselves (the firmware source, the CAD, the test output), and '
    'anything still unknown is marked "confirm".')
a.bullets([
    ('Foundations chapter: ', 'voltage, ground, pull-up resistors, the multimeter, soldering, and the software words used everywhere else.'),
    ('Computer set-up in detail: ', 'what is in the zip, installing ESP-IDF on Windows, macOS or Linux, and the desktop checks with what PASS looks like.'),
    ('Reading the board: ', 'the boot log line by line, every error message the firmware can print, the flash memory map, and the service screen.'),
    ('The wings in depth: ', 'a close-up of the switch board, how to find the right switch legs, and the timing from thumb to flap.'),
    ('The game itself: ', 'how to play, the scoring table, the HUD, and the DART trial numbers.'),
    ('How the picture is made: ', 'the renderer, the two CPU cores, and why band order matters.'),
    ('Power, print, sticker: ', 'what the firmware sets in the power chip, LiPo safety, CAD numbers, print settings, sticker printing.'),
    ('New shop sheets: ', 'E (measurement logs), F (pin and wire reference card), G (glossary), plus a FAQ and the verification data.'),
])
a.h2('How to read this manual')
a.table(['If you are\u2026', 'Read first', 'Then'], [
    ['brand new to electronics', 'Foundations (after Step 1)', 'Steps 2\u20137 in order, one gate at a time'],
    ['comfortable with Arduino / ESP32', 'Quick Start + Step 3 (computer set-up)', 'Steps 4\u20137; use Shop Sheet F as the pin card'],
    ['only printing the shell', 'Step 9 and Step 10', 'Do not close a shell around an unproven board'],
    ['curious how the game works', '"How the STRUTHIO software works"', 'Shop Sheet G, the glossary'],
    ['stuck', 'Step 12 and "Fault-finding by symptom"', 'Go back one gate (Shop Sheet D)'],
], [2.2, 2.6, 3.0], fill='CFE2F3')
a.table(['You will see', 'It means'], [
    ['Blue AUDIO GUIDE block', 'The narrated "why": read it before doing the step.'],
    ['Yellow box', 'A rule of thumb or a gate you must pass.'],
    ['Red box', 'Stop: a safety rule or something that can damage parts.'],
    ['Grey box in a fixed-width font', 'Text you type into the terminal, or text the board prints.'],
    ['"confirm"', 'Not known for certain from the project files: check it on the part in your hand.'],
    ['"example" next to a number', 'An illustration, not a measurement. Record your own.'],
], [2.2, 5.6], fill='CFE2F3')

# ------------------------------------------------------------------------------------------------
# 2. FOUNDATIONS, before Step 2
a = At(before_break(find_heading('Step 2 \u2014 Gather tools, parts, and a safe workspace')))
a.h1('FOUNDATIONS \u2014 the electronics and software you need, in plain words')
a.p('You do not need an electronics degree for this build. You need about a dozen ideas, one meter, and a soldering '
    'iron you are not afraid of. This chapter gives you exactly those, using the parts of this project as the examples. '
    'If you already know them, skim the tables and move on.')
a.audio('AUDIO GUIDE \u2014 Foundations A: Electricity for this build', [
    'Think of voltage as pressure and current as flow. The board works at 3.3 volts inside. USB supplies 5 volts, and '
    'a single lithium cell sits around 3.7 volts on average, more when full and less when empty. The board\'s power '
    'chip turns whatever arrives into the steady 3.3 volts the processor needs, so you never feed the processor directly.',
    'Ground, written GND, is the zero-volt reference that every voltage is measured from. Every circuit is a loop: '
    'current leaves a supply, does its work, and returns to ground. That is why each wing switch needs only two wires. '
    'One goes to the input pin, the other to ground.',
    'A short circuit is a loop with nothing in it to limit the current: a wire, a solder bridge or a dropped screwdriver '
    'joining a supply straight to ground. On USB it usually just makes the board reset or shut off. On a lithium cell it '
    'can make wires glow and the cell swell. That is why the battery comes last and why you check for shorts with the '
    'meter before you apply power.',
    'Polarity means which side is plus and which is minus. Resistors and switches do not care. Batteries, speakers and '
    'many connectors do, and the battery cares the most. A red wire is not proof of plus: check with the meter.',
])
a.table(['Word', 'What it means on this build', 'Where you meet it'], [
    ['Voltage (V)', '"Pressure". 3.3 V logic, 5 V USB, ~3.0\u20134.2 V lithium cell.', 'Meter test 3, battery work'],
    ['Current (A, mA)', '"Flow". The board draws a few hundred mA; measure it, do not guess.', 'Step 8, Shop Sheet E'],
    ['GND (ground)', 'The 0 V return every circuit shares.', 'Wing wiring, every measurement'],
    ['Short circuit', 'A supply joined straight to GND. Dangerous with a battery.', 'Meter test 2, before every power-up'],
    ['Polarity', 'Which lead is + and which is \u2212.', 'Battery, speaker'],
    ['Pull-up resistor', 'A weak resistor that holds an input HIGH until a switch pulls it LOW.', 'Inside the ESP32 for GPIO17/18'],
    ['Active low', '"Pressed" reads 0 (LOW), "released" reads 1 (HIGH).', 'Wing buttons, Step 6'],
    ['Normally open (NO)', 'A switch that connects only while pressed.', 'B3F wing switches'],
    ['Continuity', 'An unbroken electrical path. The meter beeps.', 'Meter test 1, switch check'],
    ['Bounce', 'A mechanical switch chatters for a few ms when it closes.', 'Removed by the 8 ms debounce'],
    ['GPIO', 'General-purpose input/output: a processor pin the firmware can read or drive.', 'GPIO17 LEFT, GPIO18 RIGHT'],
    ['I2C, SPI/QSPI, I2S', 'Wiring "languages" between chips: control, display data, audio.', 'Sheets S1 and S4 (on the board already)'],
], [1.6, 4.0, 2.2])
a.figure(FG('f01_pullup.png'), 'FIGURE F1 \u2014 Why a wing needs only two wires: the ESP32\'s internal pull-up holds GPIO17 at 3.3 V (reads 1); '
         'the closed switch pulls it to 0 V (reads 0). The firmware treats 0 as "pressed".')
a.h2('The multimeter: the only test instrument you need')
a.p('Any basic digital multimeter will do. Learn three settings and you can find almost every wiring fault in this project.')
a.figure(FG('f02_multimeter.png'), 'FIGURE F2 \u2014 Continuity (beep) for wires and switches, the same beep setting as a short check before power, '
         'and DC volts for the battery and rails. The 3.95 reading is an example.')
a.bullets([
    'Continuity and resistance only on an UNPOWERED circuit: unplug USB and the battery first.',
    'For volts, the red lead stays in the V\u03a9 socket. Never use the 10 A / mA socket for voltage: that socket is a near short.',
    'Touch the probes together once at the start: the meter should beep. Now you know the meter works.',
    'Write the reading down (Shop Sheet E). A number you did not record is a number you will measure again.',
])
a.h2('Soldering for first-timers')
a.p('Every solder joint in this build is a large, through-hole joint on perfboard or a wire on a pad. These are the '
    'easiest joints there are. Practise ten of them on scrap before you touch the switches.')
a.steps([
    ('Set up. ', 'Iron in its stand, tip wiped on a damp sponge or brass wool, 320\u2013350 \u00b0C for leaded solder or 350\u2013380 \u00b0C '
     'for lead-free (follow your solder\'s label). Fan or fume extractor on. Safety glasses for clipping leads.'),
    ('Tin the tip. ', 'Melt a little solder onto the clean tip so it shines. A shiny tip moves heat; a black one does not.'),
    ('Heat both parts. ', 'Touch the tip so it presses on the pad AND the leg (or wire) at the same time. Count two seconds.'),
    ('Feed the solder into the joint, ', 'on the side opposite the iron, not onto the iron. It should flow and wrap the leg.'),
    ('Remove solder, then the iron. ', 'Keep everything still for three seconds while it sets.'),
    ('Inspect. ', 'Compare with Figure F3. Shiny and concave is good. Fix anything else now, while it is easy.'),
    ('Pre-tin wires. ', 'Strip 3 mm, twist the strands, melt a little solder into them. A tinned wire solders to a pad in one second.'),
])
a.figure(FG('f03_solder.png'), 'FIGURE F3 \u2014 Good and bad joints from the side. A bridge on the wing board joins signal to ground: the wing then reads "pressed" all the time.')
a.callout('Lead solder and flux fumes: work ventilated, do not eat at the bench, wash your hands afterwards. '
          'The iron is 350 \u00b0C: always return it to the stand and switch it off when you leave.', fill='F4CCCC')
a.h2('Practice exercise (20 minutes)')
a.steps(['Solder five short wires into a scrap of perfboard.', 'Join pairs of them on the underside with a solder bridge on purpose, then remove it with wick.',
         'Use the meter\'s continuity beep to prove which holes are joined and which are not.',
         'Solder a spare tactile switch (any 4-leg one) and find its joined leg pairs with the meter (see Figure 6.1 in Step 6).'])
a.h2('Software words in plain language')
a.table(['Word', 'What it means', 'Example on this project'], [
    ['Host / desktop', 'Your computer.', 'Runs the checks in handheld/host'],
    ['Target / device', 'The ESP32-S3 board.', 'Runs handheld/firmware'],
    ['Terminal', 'A text window where you type commands.', 'Windows: "ESP-IDF 5.5 CMD"; macOS: Terminal'],
    ['Repository / repo', 'A folder of project files with history (git).', 'The zip is a snapshot of it'],
    ['Build / compile', 'Turn C source text into a program.', 'idf.py build'],
    ['Flash', 'Write the program into the board\'s flash memory over USB.', 'idf.py flash'],
    ['Monitor / log', 'Watch the text the board prints over USB.', 'idf.py monitor (exit: Ctrl + ])'],
    ['Firmware', 'The program that runs on the board.', 'build/struthio_a0.bin'],
    ['Partition', 'A named region of the flash chip.', 'factory (app), assets (art), nvs (settings)'],
    ['ESP-IDF', 'Espressif\'s official tools and libraries for the ESP32.', 'Version 5.5 or newer'],
    ['Tick', 'One step of the game: 60 per second.', 'The golden traces are counted in ticks'],
    ['Trace', 'A recording of inputs plus the expected state each tick.', 'golden/climb.trace'],
    ['Digest / SHA-256', 'A 64-hex-digit fingerprint of the whole game state.', 'One wrong bit changes it completely'],
    ['Asset pack', 'All pictures the device draws, in one file.', 'build/assets/struthio.pak (8.3 MB)'],
], [1.6, 3.4, 2.8])
a.callout('You cannot damage the board by building or flashing software. You CAN damage it with a short circuit, '
          'a reversed battery, or a wire touching the wrong pin. Be relaxed with commands and careful with wires.')

# ------------------------------------------------------------------------------------------------
# 3. Step 3 additions, before Step 4
a = At(before_break(find_heading('Step 4 \u2014 Prove the stock board on the bench')))
a.h2('What is in the STRUTHIO zip')
a.p('Unzip STRUTHIO_HANDHELD_v0.9.zip into a folder whose path has no spaces (for example C:\\struthio or ~/struthio). '
    'Everything below lives under handheld/.')
a.table(['Folder or file', 'What it is', 'Do you need it?'], [
    ['core/', 'The game: simulation, input normalizer, digests, replay (portable C).', 'Yes, built automatically'],
    ['render/', 'Scene builder, panel renderer, asset-pack loader, greybox renderer.', 'Yes, built automatically'],
    ['firmware/', 'The ESP-IDF app: board adapter, tasks, service mode.', 'Yes: you build and flash this'],
    ['build/assets/struthio.pak', 'The prebuilt 8.3 MB asset pack.', 'Yes: flashed by idf.py flash'],
    ['golden/', 'The five browser-recorded traces.', 'For the checks'],
    ['host/', 'Desktop programs: replay, image checks, pack builder.', 'For the checks'],
    ['cad/a1/', 'STRUTHIO084.scad, STL print files, sticker art, renders.', 'Step 9 and 10'],
    ['cad/ (A0 files)', 'The earlier bench enclosure and its renders.', 'Reference'],
    ['docs/', 'Engineering manual, pinout, bring-up checklist, schematics.', 'Reference'],
    ['tools/', 'Generators for the manuals, the C tables and the reference captures.', 'No (advanced)'],
    ['run_tests.sh', 'Every desktop check in one command.', 'Needs the full git repository (see below)'],
], [2.0, 3.8, 2.0])
a.note('You do NOT need to run "make -C handheld/host pak". The pack is already built and included. Rebuilding it '
       'needs the browser reference captures: unzip the four STRUTHIO_HANDHELD_reference_captures_part*.zip files '
       'inside handheld/ so that handheld/build/reference/ exists, then run the command.')
a.h2('Install ESP-IDF, step by step')
a.p('ESP-IDF is Espressif\'s free toolkit for the ESP32. Use version 5.5 or newer; the project\'s compile check uses '
    'v5.5.5. Follow Espressif\'s "Get Started" guide for your system. The outline is:')
a.table(['System', 'Install', 'Open a ready terminal'], [
    ['Windows 10/11', 'Download and run the "ESP-IDF Tools Installer" from Espressif; choose v5.5.x and the ESP32-S3 target.',
     'Start menu: "ESP-IDF 5.5 CMD" or "ESP-IDF 5.5 PowerShell"'],
    ['macOS', 'Install the prerequisites from the guide (Homebrew: cmake ninja dfu-util python3), then clone ESP-IDF v5.5.x and run ./install.sh esp32s3.',
     'In each new Terminal: . $HOME/esp/esp-idf/export.sh'],
    ['Linux', 'Install the prerequisites from the guide (git, wget, flex, bison, gperf, python3, cmake, ninja, ccache, libffi-dev, libssl-dev, dfu-util, libusb-1.0), clone v5.5.x, ./install.sh esp32s3.',
     'In each new terminal: . $HOME/esp/esp-idf/export.sh'],
], [1.3, 4.2, 2.3], fill='CFE2F3')
a.code(['# macOS / Linux, once:', 'mkdir -p ~/esp && cd ~/esp',
        'git clone -b v5.5.5 --recursive https://github.com/espressif/esp-idf.git',
        'cd esp-idf && ./install.sh esp32s3', '', '# every new terminal:', '. ~/esp/esp-idf/export.sh',
        'idf.py --version          # should print ESP-IDF v5.5.x'])
a.bullets([
    'Linux only: add yourself to the serial-port group once (sudo usermod -aG dialout $USER), then log out and in.',
    'Windows: if the board does not appear as a COM port, try another USB-C cable: many cables are charge-only.',
    '"idf.py: command not found" means this terminal has not run export.sh (or is not the ESP-IDF shortcut).',
])
a.h2('Your first desktop check')
a.p('Before any hardware, prove that the C game on your computer equals the browser game. You need a C compiler '
    'and make: on Windows use WSL (Ubuntu) or MSYS2; on macOS run xcode-select --install; on Linux install build-essential.')
a.code(['cd handheld', 'make -C host test             # the C game vs the browser, every tick',
        'make -C firmware/host_test run  # wing-button unit tests + DART trial report'])
a.p('What PASS looks like (your output must match these lines exactly, including the chain fingerprints):')
a.code(['PASS climb    10011 ticks  score  60250  round  4  deaths   0  jousts  13  eggs   3  darts  30  chain 3c9d8a4075eebdab',
        'PASS duel     20000 ticks  score  37750  round  3  deaths  23  jousts  37  eggs   3  darts 237  chain 88879e49392e99c1',
        'PASS late     10990 ticks  score 136250  round 11  deaths   0  jousts  60  eggs  13  darts  67  chain 66e9011ca571b5f7',
        'PASS mortal   14143 ticks  score  77500  round  5  deaths  12  jousts  23  eggs   3  darts  81  chain 25b7dd681c36acce',
        'PASS raw      20000 ticks  score  65750  round  5  deaths  30  jousts  65  eggs  15  darts 194  chain cfa154009f20e9e1',
        'GOLDEN REPLAY: all 5 traces bit-exact', '', 'PASS: wing buttons (debounce, chord, DART trials A/B/C, service)'], size=7)
a.p('The full suite, run_tests.sh, also re-checks the C against the browser game\'s own source files and pictures, so it '
    'needs the complete git repository (with arcade/), Node 22 and the reference captures. On the project machine it '
    'prints "ALL HANDHELD CHECKS PASS" in about a minute and a half. With the zip alone, the two commands above are the right check.')
a.table(['Check (run_tests.sh)', 'What it proves'], [
    ['port authority', 'The browser game the C was ported from has not changed since.'],
    ['golden traces', 'The five traces are re-recorded from the browser byte for byte.'],
    ['C core replay', 'The C game equals the browser every tick of 75,144 ticks (SHA-256 per tick).'],
    ['scene check', 'The C scene builder produces the browser\'s exact draw lists.'],
    ['raster / panel checks', 'The C pictures match the browser\'s (\u2265 30 dB at 3x scale, \u2265 24 dB on the 320 x 480 panel).'],
    ['band order', 'The two-core band hand-off never sends a band out of order (200 frames).'],
    ['wing buttons', 'Debounce, chord, DART trials and service-mode entry behave as specified.'],
    ['host shim / real headers', 'Every firmware file compiles against ESP-IDF 5.5.5\'s real headers (not yet an Xtensa link).'],
], [2.2, 5.6])
a.callout('If a desktop check fails, stop. A failure here is never a wiring problem, and the board cannot fix it.')

# ------------------------------------------------------------------------------------------------
# 4. Step 5 additions, before Step 6
a = At(before_break(find_heading('Step 6 \u2014 Build the two wing buttons')))
a.figure(FG('f11_bench.png'), 'FIGURE 5.1 \u2014 The whole A0 bench. Until the A0 pass condition is met, nothing else is connected.')
a.h2('The four commands, explained')
a.code(['cd handheld/firmware', 'idf.py set-target esp32s3      # once per fresh folder: choose the chip',
        'idf.py build                    # compile (first time: several minutes)',
        'idf.py -p PORT flash            # write bootloader, app and asset pack',
        'idf.py -p PORT monitor          # watch the log; leave with Ctrl + ]'])
a.table(['Your system', 'PORT is usually', 'How to find it'], [
    ['Windows', 'COM3, COM4, \u2026', 'Device Manager > Ports (COM & LPT): the entry that appears when you plug the board in'],
    ['macOS', '/dev/cu.usbmodem\u2026', 'ls /dev/cu.* before and after plugging in'],
    ['Linux', '/dev/ttyACM0', 'ls /dev/ttyACM* before and after plugging in'],
], [1.3, 1.8, 4.7], fill='CFE2F3')
a.bullets([
    'idf.py build flash monitor does all three in one go once PORT is known (idf.py -p PORT build flash monitor).',
    'If flashing says it cannot connect: hold BOOT, tap RESET, release BOOT, then flash again. Press RESET afterwards to run.',
    'The first flash takes longer because the 8.3 MB asset pack is written too. Later flashes rewrite it each time; that is normal.',
    'Without build/assets/struthio.pak the build prints a CMake WARNING and you get the greybox picture. That is the plan for the first boot.',
])
a.h2('Reading the boot log, line by line')
a.p('The log is the board talking to you. These are the lines the STRUTHIO firmware prints, in order. The numbers in '
    'angle brackets are filled in by the board.')
a.table(['Log line (as printed)', 'What it means', 'Good when'], [
    ['STRUTHIO Prototype A0 boot (core: STRUTHIO ARCADE 1.8.0 port)', 'Your firmware is running.', 'It appears after every reset'],
    ['AXP2101 id 0x<id>, battery present|absent', 'The power chip answered on I2C and was configured.', 'On the bench: "absent" is correct'],
    ['AXS15231B 320x480 QSPI at 40 MHz', 'The display controller was initialised.', 'Always'],
    ['DART trial <A HOLD|B TAP-HOLD|C BOTH-HOLD>, best <n>, display ok, renderer <panel|greybox>', 'Saved settings, high score, and which renderer runs.', 'renderer "greybox" without the pack, "panel" with it'],
    ['NEW RUN seed <n>', 'A game started.', 'Right after boot, and after each restart chord'],
    ['tick <n> sim <us> us (max <us>) scene <us> us render <us> us present <us> us missed <n> band-order <n>', 'Periodic performance report (see Step 7).', 'band-order 0'],
    ['ROUND <n> CLEAR score <n>', 'You cleared a round.', ''],
    ['GAME OVER score <n> (best <n>)', 'All Joust Marks used. Hold both wings to restart.', ''],
], [3.4, 2.6, 1.8])
a.h2('Error lines and what to do first')
a.table(['If the log says', 'Likely cause', 'First action'], [
    ['I2C bus failed', 'The I2C driver could not start (pins or a driver clash).', 'Re-flash the unmodified firmware; check the board revision.'],
    ['AXP2101 not found', 'The power chip did not answer at 0x34.', 'Check you have the 3.5B board; run the vendor example again.'],
    ['TCA9554 not found: the LCD stays in reset', 'The I/O expander that releases the LCD reset did not answer.', 'Board revision / I2C problem; the screen will stay dark.'],
    ['SPI bus failed / panel IO failed / AXS15231B failed', 'The display path could not be created.', 'Vendor example first; compare the board revision with the pin map.'],
    ['backlight PWM failed', 'The backlight timer could not start.', 'The picture may be there but invisible: shine a torch on the glass.'],
    ['band at y <n>: transfer timeout', 'A display transfer did not finish in time.', 'Note how often; report it with the log. Do not edit the art.'],
    ['no assets partition', 'The partition table on the board is not STRUTHIO\'s.', 'Flash with idf.py flash (it writes the table too).'],
    ['assets: mmap failed', 'The pack could not be mapped into memory.', 'Re-flash; if it persists, idf.py erase-flash then flash.'],
    ['assets: <reason> (flash build/assets/struthio.pak; see README)', 'The pack is missing or damaged.', 'Check the file exists, then flash again.'],
    ['assets: out of memory', 'Not enough RAM for the renderer\'s work buffers.', 'Report with the log; the greybox still runs.'],
], [2.6, 2.6, 2.6], fill='FCE5CD')
a.figure(FG('f06_flashmap.png'), 'FIGURE 5.2 \u2014 The flash map from firmware/partitions.csv. The high score and the DART trial live in nvs; '
         '"idf.py erase-flash" clears them.')

# ------------------------------------------------------------------------------------------------
# 5. Step 6 additions, before Step 7
a = At(before_break(find_heading('Step 7 \u2014 Bench-test the controls and play before enclosure work')))
a.h2('Build one switch board, step by step')
a.figure(FG('f04_switchboard.png'), 'FIGURE 6.1 \u2014 The wing switch board. Wire two diagonally opposite legs and prove them with the meter before soldering wires.')
a.steps([
    'Cut perfboard to about 18 x 18 mm (7 x 7 holes at 2.54 mm pitch). Sand the edges.',
    'Push the B3F-4050 through so all four legs come out underneath and the body sits flat. Bend two legs slightly to hold it.',
    'Solder all four legs (four legs = a stronger board), even though only two carry signal.',
    'Meter on beep. Find two diagonally opposite legs that are silent when released and beep when pressed (Figure 6.1, right).',
    'Mark them with a pen: S (signal) and G (ground).',
    'Pre-tin the wires. LEFT: white W1 to S, black W3 to G. RIGHT: blue W2 to S, black W4 joins the two G legs (daisy chain). Lengths: Shop Sheet F.',
    'Slide heat-shrink or a dab of hot glue over the joints for strain relief. Label the far ends LEFT / RIGHT / GND.',
    'Meter check before connecting to the board: S to G silent; press: beep. S to S of the other switch: silent.',
])
a.callout('Wing wires go to GPIO17, GPIO18 and GND only. Never to a 3V3 or 5V pin: a switch that joins a supply pin to GND '
          'is a dead short every time you press it.', fill='F4CCCC')
a.figure(FG('f05_timing.png'), 'FIGURE 6.2 \u2014 From thumb to flap. Values from firmware/main/struthio_buttons.h (1 kHz sampling, 8 ms debounce) '
         'and core/struthio_input.h (100 ms chord, 8-tick flap buffer).')
a.table(['Timing', 'Value', 'Why'], [
    ['Sampling', 'every 1 ms (1 kHz input task on core 0)', 'Fine enough to measure thumb timing'],
    ['Debounce', '8 ms stable before a change counts', 'Hides switch bounce; the press keeps its first-edge time'],
    ['Chord window', '100 ms', 'Opposite wing inside it = STRAIGHT (up)'],
    ['Flap buffer', '8 ticks (133 ms)', 'A flap that cannot fire yet is kept, not lost'],
    ['Service-mode entry', 'both wings held 650 ms inside the 800 ms boot guard', 'Hard to trigger by accident'],
    ['DART trial C', 'both wings held 200 ms', 'The default; see Step 7'],
], [1.8, 3.0, 3.0])

# ------------------------------------------------------------------------------------------------
# 6. Step 7 additions, before Step 8
a = At(before_break(find_heading('Step 8 \u2014 Add audio and choose the beginner-safe power plan')))
a.h2('The service screen, line by line')
a.figure(FG('service.png'), 'FIGURE 7.1 \u2014 The service screen, drawn on the desktop with the firmware\'s own text renderer and layout. '
         'The numbers are EXAMPLES; record what your board shows.', width=2.6)
a.table(['Line', 'Meaning', 'What you want'], [
    ['A0 / CORE 1.8.0 PORT / <date>', 'Build: when this firmware was compiled.', 'Today\'s date after you flash'],
    ['RESET <n>  PSRAM <n>K  RAM <n>K', 'Why the chip last reset (ESP-IDF reset-reason number) and free memory.', 'Write the numbers in Shop Sheet E'],
    ['LEFT UP|DOWN <n>   RIGHT UP|DOWN <n>', 'Live state of each wing and its press count.', 'DOWN only while pressed; +1 per press'],
    ['DART <trial>  FIRED <n>', 'Active DART trial and darts fired.', 'Change trial with a LEFT tap'],
    ['GOLDEN PASS <n> TICKS', 'climb.trace (10,011 ticks) replayed with every tick\'s SHA-256 checked.', 'PASS 10011 (green). FAIL is red.'],
    ['SIM <n> US/TICK  +DIGEST <n>', 'Time per game tick, and with the SHA-256 added.', 'SIM far below 16,667 us (one tick at 60 Hz)'],
    ['PANEL <n> MS/FRAME <n> FPS', 'Time to send a full 320 x 480 picture.', 'Record it; under 33 ms allows 30 fps'],
    ['LEFT: NEXT DART TRIAL \u2026', 'LEFT tap: next trial (saved). RIGHT tap: run the checks again.', 'Power-cycle to play'],
], [2.4, 3.4, 2.0])
a.p('Both timings include a short pause every 256 ticks so the watchdog stays happy, so they read slightly high. '
    'They are the project\'s first real device measurements: they are still unknown until your board prints them.', italic=True)
a.h2('The DART trials, by the numbers')
a.p('DART has no button of its own. Three ways to ask for it from two wings were tested on the desktop by feeding the '
    'golden bots\' recorded button timelines through each trial. "Fired" counts darts the trial would have fired; the bots '
    'were steering and flapping, not trying to dart, so almost every fired dart is an accident.')
a.table(['Trace (length)', 'A HOLD (230 ms one wing)', 'B TAP-HOLD (re-press within 220 ms, hold 60)', 'C BOTH-HOLD (both 200 ms)'], [
    ['climb (2.8 min)', '194 fired (69.8/min)', '110 (39.6/min)', '1 (0.4/min)'],
    ['duel (5.6 min)', '390 (70.2/min)', '232 (41.8/min)', '4 (0.7/min)'],
    ['late (3.1 min)', '212 (69.4/min)', '149 (48.8/min)', '1 (0.3/min)'],
    ['mortal (3.9 min)', '308 (78.4/min)', '184 (46.8/min)', '8 (2.0/min)'],
], [1.6, 2.0, 2.4, 1.8], fill='CFE2F3')
a.p('Trial A and B misfire constantly because steering IS holding a wing, and flapping while steering IS a quick re-press. '
    'Trial C uses the one gesture normal play leaves free. That is why C is the default, and why the final choice still '
    'belongs to real thumbs (Chapter 9).')
a.h2('How to play STRUTHIO')
a.bullets([
    ('Flap: ', 'tap LEFT to flap up and to the left, RIGHT to flap up and to the right.'),
    ('Straight up: ', 'press both wings within 100 ms of each other.'),
    ('Steer: ', 'hold a wing.'),
    ('Joust: ', 'meet a rival; the higher lance wins, and nearly level is a clash where both bounce. A beaten rival falls as a rider and turns into an egg when it lands.'),
    ('Eggs: ', 'collect them before they hatch: the rulebook egg lasts 360 ticks (6 s), hatches for 60 ticks (1 s), then the rival remounts in 90 ticks (1.5 s).'),
    ('Rings: ', 'fly through all six rings of the tower in any order. Then the gold ring at the moon opens; take it to blast every rival and clear the round.'),
    ('Lava: ', 'touch it and you start to sink. Flap within 18 ticks (0.3 s) to escape, or you lose a Joust Mark (a life).'),
    ('Restart: ', 'after GAME OVER, hold both wings together. A single flap cannot restart, so you will not skip the score by accident.'),
])
a.table(['Scoring event', 'Points (from the rulebook, arcade/src/data/rules.mjs)'], [
    ['Joust win', 'BOUNDER 500, HUNTER 750, SHADOW 1,000, plus 250 per tier'],
    ['Eggs in a row', '250, 500, 750, then 1,000 each (the chain resets when you lose a life)'],
    ['Ring', '500'],
    ['Gold-ring blast', 'each armed rival scores as a joust win'],
    ['Round clear', '5,000 + 1,000 per round, up to 15,000'],
    ['Survival bonus', '3,000 when you clear a round without losing a life'],
    ['Joust Marks (lives)', 'start with 11 (also the maximum); extra at 30,000, then every 100,000'],
], [2.2, 5.6])
a.figure(FG('f09_hud.png'), 'FIGURE 7.2 \u2014 The game screen as the C panel renderer draws it. The HUD fields are the browser game\'s.', width=5.6)
a.h2('How the picture is made')
a.figure(FG('f10_pipeline.png'), 'FIGURE 7.3 \u2014 From button to glass. The same C files run in the desktop checks.')
a.figure(R('docs', 'renders', 'panel_vs_browser_2809.png'), 'FIGURE 7.4 \u2014 The C panel renderer (left) and the browser (right), the same tick. '
         'Measured closeness: 24.9\u201330.3 dB PSNR over 18 frames; 30 dB and above looks identical to most eyes.', width=4.4)
a.figure(R('docs', 'renders', 'a0_greybox_sheet.png'), 'FIGURE 7.5 \u2014 The greybox renderer: what the first boot without the asset pack looks like. Same game, flat colours.', width=5.0)
a.figure(FG('f07_cores.png'), 'FIGURE 7.6 \u2014 The two cores and the band order. A band-order count above 0 in the log means the hand-off is broken.')
a.h2('The performance line, decoded')
a.code(['tick 3600 sim 180 us (max 410) scene 2100 us render 9800 us present 7400 us missed 0 band-order 0   <- EXAMPLE numbers'])
a.table(['Field', 'Meaning', 'Good'], [
    ['tick', 'Game ticks since boot (60 per second).', 'Rises by 60 per second'],
    ['sim / max', 'Time for one game step, average and worst.', 'Far below 16,667 us'],
    ['scene', 'Time to build the list of textured quads.', 'Record it'],
    ['render', 'Time for the panel renderer to draw the frame.', 'Record it'],
    ['present', 'Time to send it to the display.', 'Record it'],
    ['missed', 'Game deadlines missed (the game fell behind 60 Hz).', '0, or very rare'],
    ['band-order', 'Bands that reached the panel out of order.', 'Always 0'],
], [1.4, 4.0, 2.4])
a.p('If render + present is above about 33 ms, the screen runs below 30 fps. The game still runs at 60 Hz; only pictures '
    'are dropped. Chapter 11 gives the investigation order.')

# ------------------------------------------------------------------------------------------------
# 7. Step 8 additions, before Step 9
a = At(before_break(find_heading('Step 9 \u2014 Print and prove the enclosure in stages')))
a.h2('Audio: what to expect from the current firmware')
a.callout('Honest status: the STRUTHIO firmware does not play game sound yet. The ES8311 audio driver is still a stub '
          '(firmware/README.md). Use Waveshare\'s own audio example to prove the speaker, codec and amplifier now; '
          'game sound arrives with a later firmware.')
a.table(['Part', 'What it does', 'How it is connected'], [
    ['ES8311', 'Audio codec: turns digital samples into an analogue signal.', 'I2C (SDA 8, SCL 7) for settings; I2S MCLK 44, BCLK 13, LRCK 15, DOUT 16, DIN 14'],
    ['NS4150B', 'Mono class-D amplifier that drives the speaker.', 'On the board, after the codec'],
    ['Speaker', '8 \u03a9, about 1 W, about 28 mm.', 'Board speaker header (position: confirm on the vendor drawing)'],
], [1.4, 3.0, 3.4])
a.h2('What the firmware sets in the power chip')
a.p('The AXP2101 power chip charges the battery, makes every supply rail, and owns the power button. STRUTHIO copies '
    'Waveshare\'s own settings exactly (firmware/main/board_pmu.cpp). The ones that matter for choosing a battery:')
a.table(['Setting', 'Value', 'What it means for you'], [
    ['Charge current', '200 mA', 'Gentle. A cell must accept at least this; check its datasheet.'],
    ['Charge target voltage', '4.1 V', 'Slightly below the common 4.2 V full charge: kinder to the cell, a little less runtime.'],
    ['Pre-charge / end-of-charge current', '50 mA / 25 mA', 'Standard lithium charging steps.'],
    ['USB input limit', '4.36 V min, 1,500 mA max', 'Needs a decent USB supply; a weak port may not charge while playing.'],
    ['System shut-down voltage', '2.6 V', 'The board turns itself off before the cell is too empty. The cell\'s own protection is the second guard.'],
    ['Power key', 'hold 4 s = off, 128 ms = on', 'The board\'s own momentary key; the slide switch is a separate decision (Chapter 15).'],
    ['Battery temperature (TS) pin', 'not measured', 'Use a cell with its own protection circuit.'],
], [2.2, 1.8, 3.8])
a.h2('LiPo safety rules')
a.table(['Always', 'Never'], [
    ['Use a protected single-cell (1S) LiPo or Li-ion.', 'Use an unprotected cell, a 2S pack, or AA / NiMH cells on the lithium connector.'],
    ['Check polarity with the meter before the first plug-in.', 'Trust wire colours or "it fits" for polarity.'],
    ['Insulate bare leads with tape the moment you cut or strip them.', 'Cut both battery leads at once with the same cutters (instant short).'],
    ['Charge on a non-flammable surface while you are there.', 'Charge a puffy, dented or hot cell. Retire it.'],
    ['Leave room in the shell: the cell must not be squeezed.', 'Squeeze, puncture, bend or solder directly onto the cell body.'],
    ['Switch the + lead only (option A).', 'Switch the \u2212 lead.'],
], [3.9, 3.9], fill='F4CCCC')
a.h2('Measuring current draw (before choosing a battery)')
a.steps([
    'Put an inline USB-C power meter (the small kind with a display) between the USB charger and the board, battery disconnected.',
    'Record the current at: boot, service screen, normal play, and play in a busy scene. Note the backlight level if you change it.',
    'Take the highest steady play value as your design current.',
    'Estimate runtime: hours \u2248 capacity (mAh) \u00d7 0.8 \u00f7 design current (mA). The 0.8 leaves margin for conversion losses and an older cell.',
    'Choose the largest protected cell that fits the 36 x 52 x 6.2 mm cavity, then measure real runtime from full to shutdown (Shop Sheet E).',
])
a.p('Example arithmetic only, not a measurement: a 1,000 mAh cell at 250 mA gives about 1,000 x 0.8 / 250 = 3.2 hours.', italic=True)

# ------------------------------------------------------------------------------------------------
# 8. Step 9 additions, before Step 10
a = At(before_break(find_heading('Step 10 \u2014 Prepare the lens, sticker art, and front-face presentation')))
a.figure(R('cad', 'a1', 'renders', 'a084_struthio_front.png'), 'FIGURE 9.1 \u2014 A0.8.4, the player\'s view: navy shell, gold wing caps, sticker art. Print file: cad/a1/STRUTHIO084.scad.', width=3.6)
a.h2('The CAD numbers that matter')
a.p('All dimensions are in millimetres and come from STRUTHIO084.scad. Change a number at the top of the file, export '
    'again, and run the checks. Do not edit the STL files themselves.')
a.table(['Parameter', 'Value', 'What it controls'], [
    ['BODY_W x BODY_H', '88 x 128', 'Outer footprint'],
    ['BASE / SPEAKER / BATTERY_TOTAL_D', '18.6 / 22.5 / 23.0', 'Depth: board-only areas, over the speaker, over the battery'],
    ['WALL / FRONT_T / BACK_WALL', '2.4 / 3.0 / 2.2', 'Side wall, front face and rear skin thickness'],
    ['FIT', '0.35', 'Clearance added around the board'],
    ['BOARD_W x BOARD_H x BOARD_D', '61.00 x 92.44 x 11.50', 'The Waveshare board envelope'],
    ['LCD_W x LCD_H', '48.96 x 73.44', 'Visible display window'],
    ['LENS_LAND_W x H, recess', '56.5 x 81.0, 0.55 deep', 'Where the clear lens sits'],
    ['BTN_X, BTN_Y', '\u00b121.5, \u221247.5', 'Wing switch centres'],
    ['BTN_FACE_W x H, protrusion', '28 x 18.5, 1.8 proud', 'Wing cap size and how far it stands out'],
    ['BTN_PRETRAVEL_GAP', '0.10', 'Gap between cap stem and switch plunger at rest'],
    ['SWITCH_PCB_PLANE_Z', '13.7', 'Height of the switch boards: the one to tune with a real cap'],
    ['SWITCH_H / SWITCH_TRAVEL', '7.3 / 0.25', 'B3F-4050 height and operating travel'],
    ['SPKR_D x SPKR_T, SPKR_Y', '28 x 5, \u221243.5', 'Speaker pocket (position provisional)'],
    ['BAT_W x BAT_H x BAT_T', '36 x 52 x 6.2', 'Battery cavity'],
    ['SCREW_X, SCREW_Y, SCREW_D', '\u00b137, \u00b155.5, 2.2', 'Four rear M2-class screws'],
    ['STICKER_TOP / BOTTOM, widths', '\u221227.2 / \u221260.0, 72 / 78', 'Art panel (sits 0.5 mm below the lens land)'],
], [2.6, 2.0, 3.2])
a.h2('Making the print files and checking them')
a.code(['cd handheld/cad/a1', './export_a1.sh            # needs OpenSCAD 2021.01 or newer',
        '#  -> stl/struthio_a084_front.stl, _back.stl, _left_button.stl, _right_button.stl',
        '#  -> svg/struthio_a084_front_sticker.svg (sticker cut template), renders/, A1 CHECK'])
a.p('The STL files in the zip are already exported, so you only need OpenSCAD if you change a parameter. '
    'check_a1.py then re-measures the geometry. On A0.8.4 all 20 checks pass, for example:')
a.code(['PASS  front: one watertight solid',
        'PASS  USB-C opening clear through the bottom wall',
        'PASS  left_button: stem rests 0.10 mm before the plunger (must be 0 < gap < 0.25)',
        'PASS  left_button: neck passes the opening with >= 0.2 mm clearance all round',
        'PASS  left_button: flange overlaps the opening by >= 0.3 mm all round (captured)',
        'PASS  left_button: 2.67 mm of shell between the opening and the speaker grille (>= 1.5)',
        'PASS  art sticker top at y -27.20 stays below the lens land (-26.72)',
        'A1 CHECK: all pass'], size=7.5)
a.h2('Print settings: a starting point')
a.p('These are common starting values for FDM printers, not project measurements. Your printer and filament may need others.', italic=True)
a.table(['Part', 'Material', 'Starting settings', 'Notes'], [
    ['Fit coupon / first front', 'PLA', '0.2 mm layers, 3 walls, 20% infill', 'Cheap and fast: for fit only'],
    ['Final front and rear', 'PETG or ASA', '0.2 mm layers, 4 walls, 25\u201330% infill', 'Front face down gives the smoothest face; check the slicer preview for supports'],
    ['Wing caps', 'PETG or ASA', '0.12\u20130.16 mm layers, 100% or 4+ walls', 'Small parts: slow down, add a brim if they lift'],
], [1.8, 1.3, 2.4, 2.3])
a.h2('Tuning the wing feel with SWITCH_PCB_PLANE_Z')
a.bullets([
    'The cap must rest just clear of the plunger: the design gap is 0.10 mm, and check_a1 requires 0 to 0.25 mm (the switch\'s whole travel).',
    'If the cap is pre-pressing the switch, or does nothing until pushed hard, change SWITCH_PCB_PLANE_Z by 0.1 mm, export, and read the new gap in check_a1.',
    'Change one number at a time. Write each value and the result in Shop Sheet E.',
])
a.figure(R('cad', 'a1', 'renders', 'wings.png'), 'FIGURE 9.2 \u2014 The wing caps and their capture flanges.', width=4.6)
a.figure(R('cad', 'a1', 'renders', 'rear_angle.png'), 'FIGURE 9.3 \u2014 The rear shell: four screw posts, debossed mark, battery and speaker depth.', width=4.6)

# ------------------------------------------------------------------------------------------------
# 9. Step 10 additions, before Step 11
a = At(before_break(find_heading('Step 11 \u2014 Final assembly from beginning to end')))
a.h2('Printing the sticker, step by step')
a.table(['File (cad/a1/)', 'Use'], [
    ['art/sticker_front_print.png', 'The print file: 600 dpi, 1919 x 850 px = 81.2 x 36.0 mm including bleed.'],
    ['art/sticker_front_proof.png', 'Proof with the cut lines drawn on, for checking on screen.'],
    ['svg/struthio_a084_front_sticker.svg', 'The exact cut outline exported from the CAD, for a cutting machine.'],
], [3.0, 4.8], fill='CFE2F3')
a.steps([
    'Print on glossy or matte sticker paper at 100% / "actual size". Turn off "fit to page" and any scaling.',
    'Measure the printed art with a ruler: it must be 81.2 mm wide. If it is not, fix the printer scaling before going on.',
    'Cover with clear laminating film (cold laminate or clear packing tape works for a prototype).',
    'Cut the outer edge and the six holes (two wings, four grille slots) with a fresh craft blade and a metal ruler, or a cutting machine using the SVG.',
    'Dry-fit on the printed front before peeling: the wing caps and grille slots must line up with clearance.',
    'Peel a corner, align the top edge first, then smooth downwards from the centre.',
])
a.figure(R('cad', 'a1', 'renders', 'a084_identity_details.png'), 'FIGURE 10.1 \u2014 A0.8.4 identity details: sticker panel, wing caps, grille and rear mark.')

# ------------------------------------------------------------------------------------------------
# 10. Step 12 additions, before Shop Sheet A
a = At(before_break(find_heading('SHOP SHEET A \u2014 BOM: buy now vs defer')))
a.h2('Fault-finding by symptom')
a.table(['Symptom', 'Check in this order'], [
    ['Nothing on screen and no log', 'Cable (data, not charge-only) \u2192 another USB port \u2192 vendor example \u2192 BOOT + RESET flash'],
    ['Log runs, screen dark', 'Backlight (torch on the glass) \u2192 "TCA9554 not found" in the log \u2192 board revision'],
    ['Picture in horizontal stripes or shifted blocks', 'band-order count in the log \u2192 "transfer timeout" lines \u2192 report with the log'],
    ['Greybox instead of the real art', 'renderer "greybox" in the boot line \u2192 "assets:" errors \u2192 struthio.pak present when you built?'],
    ['GOLDEN FAIL on the service screen', 'Copy the whole log line (it names the tick and the reason) \u2192 does make -C host test pass on the desktop?'],
    ['A wing never registers', 'Service screen state \u2192 meter: switch legs (diagonal?) \u2192 wire to the right header pin \u2192 GND joined'],
    ['A wing is always "DOWN"', 'Solder bridge on the switch board \u2192 wrong leg pair (joined inside) \u2192 wire touching GND'],
    ['Double flaps from one press', 'Press count in service mode rises by 2? \u2192 cracked or cold joint \u2192 loose wire'],
    ['Wings fine until the screen turns on', 'A camera module fitted? (GPIO17/18 are its VSYNC/HREF) \u2192 board revision'],
    ['Resets during play', 'Log reason after reset (watchdog or brown-out) \u2192 USB supply \u2192 newly added part (speaker, battery) removed?'],
    ['High score lost', 'Was idf.py erase-flash run? (it clears nvs) \u2192 otherwise report with the log'],
    ['Charges on USB, dead on battery', 'Polarity at the connector \u2192 slide switch position and wiring \u2192 cell protection tripped (charge it on USB)'],
], [2.6, 5.2], fill='FCE5CD')

# ------------------------------------------------------------------------------------------------
# 11. NEW chapter: how the software works, before Shop Sheet A
a.h1('How the STRUTHIO software works (for the curious)')
a.p('You can build the whole handheld without reading this chapter. Read it when you want to know why the project trusts '
    'a board it has never seen.')
a.audio('AUDIO GUIDE \u2014 Software A: One game, proven twice', [
    'STRUTHIO began as a browser game, STRUTHIO ARCADE 1.8.0, written in JavaScript. A small handheld cannot run a browser, '
    'so the game was rewritten in portable C. A rewrite is easy to get almost right and very hard to get exactly right. '
    'Almost right is not good enough: a tiny difference in one bounce changes everything that happens after it.',
    'So the project records "golden traces" from the browser. A trace stores the inputs for every tick and a SHA-256 '
    'fingerprint of the whole game state after every tick. The C game replays the same inputs and must produce the same '
    'fingerprint, tick after tick. Five traces cover 75,144 ticks, about 21 minutes of play, including deaths, rounds up to 11 and hundreds of jousts. '
    'All five match exactly.',
    'The firmware carries one of those traces inside it. Service mode replays it on the ESP32 itself. If the screen says '
    'GOLDEN PASS, the chip in your hands is running the same game, not an imitation of it.',
])
a.figure(FG('f08_golden.png'), 'FIGURE S1 \u2014 The golden replay: the same trace proves the C game on the desktop and on the device.')
a.table(['Trace', 'Ticks', 'Score', 'Round', 'Deaths', 'Jousts', 'Eggs', 'Darts'], [
    ['climb', '10,011', '60,250', '4', '0', '13', '3', '30'],
    ['duel', '20,000', '37,750', '3', '23', '37', '3', '237'],
    ['late', '10,990', '136,250', '11', '0', '60', '13', '67'],
    ['mortal', '14,143', '77,500', '5', '12', '23', '3', '81'],
    ['raw', '20,000', '65,750', '5', '30', '65', '15', '194'],
], [1.2, 1, 1, 0.8, 0.8, 0.8, 0.8, 0.8], fill='CFE2F3')
a.h2('Map of the source code')
a.table(['File', 'Job'], [
    ['core/struthio_sim.c', 'The game rules and physics, one tick at a time (st_step).'],
    ['core/struthio_tower.c, struthio_rules.h', 'Tower layout and rulebook values, generated from the browser game.'],
    ['core/struthio_input.c', 'The input normalizer: wing presses \u2192 flaps, STRAIGHT chord, steering, DART.'],
    ['core/struthio_digest.c, struthio_replay.c', 'SHA-256 of the state; golden-trace replay.'],
    ['render/struthio_scene.c', 'Builds the list of textured quads for a frame, exactly as the browser does.'],
    ['render/struthio_panel.c', 'Draws those quads in 16-row bands at 320 x 480 for the panel.'],
    ['render/struthio_pak.c', 'Reads the asset pack mapped from flash.'],
    ['render/struthio_greybox.c', 'The flat-colour fallback renderer and the service-screen text.'],
    ['firmware/main/main.c', 'The tasks: input (1 kHz), game (60 Hz), render (two cores), NVS saving.'],
    ['firmware/main/struthio_buttons.c', 'Debounce, press counting, DART trials, service-mode entry.'],
    ['firmware/main/service.c', 'The service screen and the on-device golden replay.'],
    ['firmware/main/board_waveshare_35b.c, board_pmu.cpp', 'Board adapter ported from Waveshare\'s example: I2C, power chip, LCD reset, QSPI panel, backlight.'],
    ['firmware/main/board_pins.h', 'The full pin map (Shop Sheet F).'],
], [3.2, 4.6])
a.h2('What is not finished yet')
a.bullets([
    'The firmware has been compiled against the real ESP-IDF 5.5.5 headers, but the first full Xtensa build and link happens on your computer.',
    'Device timings (simulation, render, panel) are unmeasured until a board prints them.',
    'Game audio: the ES8311 driver is a stub.',
    'DART gesture: trial C is the default; the decision belongs to real thumbs.',
    'Battery capacity, slide-switch option, USB-C extension part: chosen after measurement.',
    'To confirm on the board in hand: expansion-header pin numbers, speaker and battery connector positions, the ES8311 address.',
])

# ------------------------------------------------------------------------------------------------
# 12. New shop sheets + glossary + FAQ, before the Appendix
a = At(before_break(find_heading('Appendix \u2014 Where the current project stands')))
a.h1('SHOP SHEET E \u2014 Measurement log')
a.p('Print this page. Fill it in as you go. These numbers are the project\'s real data: nobody has them yet but you.')
a.h2('Board and firmware')
a.table(['Item', 'Your value'], [['Board revision printed on the PCB', ''], ['ESP-IDF version (idf.py --version)', ''],
                                ['Firmware build date (service screen)', ''], ['Serial port name', ''], ['AXP2101 id (boot log)', '']], [3.6, 4.2])
a.h2('Service screen')
a.table(['Run', 'RESET', 'PSRAM K', 'RAM K', 'GOLDEN', 'SIM us', '+DIGEST us', 'PANEL ms', 'FPS'],
        [[str(i), '', '', '', '', '', '', '', ''] for i in (1, 2, 3)], [0.5, 0.8, 0.9, 0.8, 1.0, 0.8, 1.0, 1.0, 0.8])
a.h2('Performance line during play')
a.table(['Scene', 'sim / max us', 'scene us', 'render us', 'present us', 'missed', 'band-order'],
        [['Attract / start', '', '', '', '', '', ''], ['Normal play', '', '', '', '', '', ''], ['Busy scene', '', '', '', '', '', ''], ['After 10 min', '', '', '', '', '', '']],
        [1.6, 1.1, 1.0, 1.0, 1.0, 0.9, 1.0])
a.h2('Current draw and battery')
a.table(['Condition', 'Current mA', 'Notes'], [['Boot', '', ''], ['Service screen', '', ''], ['Normal play', '', ''],
                                           ['Busy play', '', ''], ['Charging (battery fitted)', '', ''], ['Battery: capacity / runtime full \u2192 off', '', '']], [3.0, 1.6, 3.2])
a.h2('Print and fit')
a.table(['Print #', 'Part', 'SWITCH_PCB_PLANE_Z', 'Rest gap (check_a1)', 'Feel / fit result'],
        [[str(i), '', '', '', ''] for i in (1, 2, 3, 4)], [0.8, 1.6, 1.6, 1.6, 2.2])

a.h1('SHOP SHEET F \u2014 Pin and wire reference card')
a.p('From firmware/main/board_pins.h and Waveshare\'s example (commit 840daf2). Bold rows are the only pins you wire.')
a.table(['Function', 'GPIO pins', 'Notes'], [
    ['LEFT WING', '17', 'Input, pull-up, active low. Header pin 16 (confirm).'],
    ['RIGHT WING', '18', 'Input, pull-up, active low. Header pin 18 (confirm).'],
    ['GND', '\u2014', 'Any GND pin, e.g. header pin 30 (confirm).'],
    ['LCD (AXS15231B, QSPI)', 'CS 12, SCLK 5, D0\u2013D3 1\u20134', 'On the board'],
    ['Backlight', '6 (PWM)', 'On the board'],
    ['LCD reset', 'TCA9554 output EXIO1', 'Via the I/O expander'],
    ['I2C', 'SDA 8, SCL 7', 'Power chip, expander, codec, touch, IMU, clock'],
    ['I2S audio', 'MCLK 44, BCLK 13, LRCK 15, DOUT 16, DIN 14', 'ES8311 codec'],
    ['SD card', 'CMD 10, CLK 11, D0 9', 'Unused by STRUTHIO'],
    ['Camera (do not fit)', 'XCLK 38, PCLK 41, VSYNC 17, HREF 18, data 45 47 48 46 42 40 39 21', 'Shares 17/18 with the wings'],
    ['BOOT / USB / UART0 TX', '0 / 19, 20 / 43', 'Do not use'],
    ['Spare (no camera)', '21, 38, 39, 40, 41, 42, 47, 48', 'For a slide-switch signal or later buttons; avoid 45/46 (strapping)'],
], [2.0, 3.0, 2.8], bold_first=True)
a.table(['I2C device', 'Address', 'Job'], [
    ['AXP2101', '0x34', 'Power, charging, power key'], ['TCA9554', '0x20', 'I/O expander (LCD reset)'],
    ['Touch controller', '0x3B', 'Unused by STRUTHIO'], ['QMI8658', '0x6B', 'Motion sensor, unused'],
    ['PCF85063', '0x51', 'Real-time clock, unused'], ['ES8311', 'confirm', 'Audio codec'],
], [2.2, 1.4, 4.2], fill='CFE2F3')
a.table(['#', 'From', 'To', 'Wire', 'Notes'], [
    ['W1', 'Header pin 16 (GPIO17)', 'LEFT switch, signal leg', 'white, ~22 AWG, ~110 mm', 'confirm the pin'],
    ['W2', 'Header pin 18 (GPIO18)', 'RIGHT switch, signal leg', 'blue, ~22 AWG, ~110 mm', ''],
    ['W3', 'Header pin 30 (GND)', 'LEFT switch, ground leg', 'black, ~22 AWG', 'any GND pin'],
    ['W4', 'LEFT ground leg', 'RIGHT ground leg', 'black, ~22 AWG, ~45 mm', 'GND daisy-chain'],
    ['W5', 'Speaker + / \u2212', 'Board speaker header', '2-core speaker lead', 'position: confirm'],
    ['W6', 'LiPo + (protected)', 'Slide switch common', 'red, rated \u2265 2 A', 'option A only'],
    ['W7', 'Slide switch ON', 'Board BAT +', 'red', 'check polarity first'],
    ['W8', 'LiPo \u2212', 'Board BAT \u2212', 'black', 'never switched'],
    ['W9', 'USB-C panel jack', 'Board USB-C', 'full-data extension', 'data needed for flashing'],
    ['W10', 'Spare GPIO (e.g. 21)', 'Slide switch \u2192 GND', 'thin 2-core', 'option B only; needs firmware'],
], [0.5, 1.9, 1.9, 1.9, 1.6], fill='FCE5CD', size=9)

a.h1('SHOP SHEET G \u2014 Glossary')
a.table(['Term', 'Meaning'], [
    ['A0 / A1', 'A0: the bench prototype that proves the electronics. A1: the finished-looking STRUTHIO shell.'],
    ['A0.8.4', 'The current A1 CAD version (STRUTHIO084.scad), the one to print.'],
    ['Active low', 'A signal that means "on" when it reads 0 V.'],
    ['Asset pack (struthio.pak)', 'All textures and HUD layers at panel resolution, 8.3 MB, written to the "assets" partition.'],
    ['AXP2101', 'The power-management chip: charger, supply rails, power key.'],
    ['AXS15231B', 'The display controller chip behind the 3.5-inch 320 x 480 panel.'],
    ['Band', '16 rows of the screen, drawn and sent as one piece. 30 bands make a frame.'],
    ['B3F-4050 / -4055', 'Omron 12 x 12 mm tactile switches; 4055 is the firmer one.'],
    ['Bounce / debounce', 'A switch chatters when it closes; the firmware waits 8 ms of stable signal.'],
    ['Chord', 'Both wings within 100 ms: STRAIGHT (fly straight up).'],
    ['DART', 'The forward dash attack; on the handheld, trial C (both wings held 200 ms).'],
    ['dB / PSNR', 'A measure of how close two pictures are; higher is closer. 30 dB looks identical to most eyes.'],
    ['Digest / SHA-256', 'A fingerprint of the full game state; one changed bit changes it completely.'],
    ['ES8311 / NS4150B', 'Audio codec / speaker amplifier on the board.'],
    ['ESP32-S3', 'The dual-core 240 MHz processor on the board (with 8 MB PSRAM and 16 MB flash).'],
    ['ESP-IDF / idf.py', 'Espressif\'s toolkit / its command-line tool.'],
    ['Flash (noun / verb)', 'The 16 MB storage chip / writing a program into it.'],
    ['Golden trace', 'A browser recording of inputs plus the expected state digest every tick.'],
    ['Greybox', 'The simple flat-colour renderer used when the asset pack is absent.'],
    ['GPIO', 'A processor pin the firmware can read or drive.'],
    ['HUD', 'The score, round, rings, rivals and Joust Marks bar at the top of the screen.'],
    ['I2C / I2S / QSPI', 'Chip-to-chip connections: control / audio / fast display data.'],
    ['Joust Marks', 'Your lives: 11 at the start.'],
    ['LiPo (1S, protected)', 'A single lithium-polymer cell with its own protection circuit.'],
    ['Monitor', 'The serial log window (idf.py monitor). Leave with Ctrl + ].'],
    ['NVS', 'A small flash area for settings: high score and DART trial.'],
    ['Partition', 'A named region of flash: factory (app), assets, nvs, phy.'],
    ['PSRAM', 'Extra 8 MB of RAM on the ESP32-S3 module.'],
    ['Pull-up', 'A resistor that holds an input high until something pulls it low.'],
    ['RAMWR / RAMWRC', 'Display commands: start a frame / continue where the last write ended.'],
    ['Service mode', 'The hidden diagnostic screen: hold both wings at power-on.'],
    ['TCA9554', 'An I/O expander chip; its output EXIO1 releases the LCD reset.'],
    ['Tick', 'One game step; 60 per second.'],
    ['Watchdog', 'A timer that resets the chip if the game hangs (3 s), returning it to play.'],
], [2.1, 5.7])

a.h1('Frequently asked questions')
faq = [
    ('Do I need to know C to build this?', 'No. You build and flash the code exactly as it is. You only type the commands in this manual.'),
    ('Can I use a different ESP32 board?', 'Not without changing the board adapter. The firmware targets the Waveshare ESP32-S3-Touch-LCD-3.5B: its display, power chip and pins.'),
    ('Why is the touch screen not used?', 'The game is played with two wings by design, like a 1990s handheld. The touch controller is on the board but unused.'),
    ('Why can\'t I fit the camera?', 'Its VSYNC and HREF lines are GPIO17 and GPIO18, the two wing inputs. Fitting it would fight the buttons.'),
    ('The first boot shows plain shapes, not the art. Is it broken?', 'No. Without struthio.pak the firmware runs the greybox renderer on purpose. Flash with the pack for the real art.'),
    ('Do I need Node or the browser game?', 'Not to build the handheld. They are only needed for the full run_tests.sh and for rebuilding the asset pack.'),
    ('Will flashing erase my high score?', 'A normal idf.py flash keeps it (nvs is not rewritten). idf.py erase-flash clears it.'),
    ('How do I get to service mode again?', 'Power off, hold both wings, power on, keep holding for about a second.'),
    ('Why is there no sound?', 'The firmware\'s audio driver is not written yet. Prove the speaker with Waveshare\'s audio example for now.'),
    ('What battery should I buy?', 'None yet. Measure current first (Step 8), then choose a protected 1S cell that fits 36 x 52 x 6.2 mm.'),
    ('Which DART trial is best?', 'Trial C misfires least in recorded play (0.3\u20132.0 a minute versus 40\u201378 for A and B). Your thumbs decide.'),
    ('The log scrolls too fast to read.', 'Press Ctrl + T then Ctrl + Y in idf.py monitor to pause output, or copy the whole window into a text file.'),
    ('Something smells hot.', 'Unplug USB and the battery at once. Do not touch the board until it is cool. Look for a short with the meter (power off).'),
]
for q, ans in faq:
    a.p(ans, bold_lead=q + '  ')

# ------------------------------------------------------------------------------------------------
# 13. Appendix additions, at the end of the document (before sectPr)
a = At(body.find(qn('w:sectPr')))
a.h2('Verification data behind this edition')
a.table(['Check', 'Result'], [
    ['C game vs browser', '5 traces, 75,144 ticks, SHA-256 equal every tick'],
    ['Scene builder vs browser', 'Draw lists equal on every tick of all 5 traces'],
    ['C rasterizer vs WebGPU (3x canvas)', '35.7\u201350.6 dB'],
    ['Panel renderer vs browser at 320 x 480', '24.9\u201330.3 dB over 18 frames (bilinear reference 28.2\u201331.2 dB)'],
    ['HUD layers', 'recompose the browser HUD at 43.2 dB'],
    ['Band order (two cores)', 'old scheme 4,736 of 6,000 bands misplaced; current scheme 0'],
    ['Desktop frame time (not the ESP32)', '8\u201310 ms per frame on the development computer'],
    ['Firmware compile', 'all sources against ESP-IDF 5.5.5 headers and board drivers, -Wall -Wextra -Werror'],
    ['A1 CAD (A0.8.4)', '20 of 20 geometry checks pass'],
], [3.0, 4.8])
a.h2('Edition history')
a.table(['Edition', 'What changed'], [
    ['Engineering v0.4\u2013v0.5', 'A0 bench package; the game ported to portable C and proven bit-exact.'],
    ['v0.6\u2013v0.7', 'Manual for the C port; hardware chapter in the 1990s handheld style.'],
    ['v0.8', 'A1 shell becomes STRUTHIO: A0.8.4 identity, wing caps, sticker art.'],
    ['v0.9', 'ESP32 preparation: board adapter from Waveshare\'s example, real-header compile check, band-order fix.'],
    ['v0.10', 'Wiring schematics S1\u2013S5.'],
    ['Build Manual 1.1', 'Beginner build order with audio guide, photos, shop sheets A\u2013D.'],
    ['Build Manual 1.2', 'This edition: foundations, set-up, logs, timing, game guide, power, print, shop sheets E\u2013G, FAQ.'],
], [2.0, 5.8])

doc.save(OUT)
print('wrote', OUT)
