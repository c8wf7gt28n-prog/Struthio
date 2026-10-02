"""STRUTHIO HANDHELD · helpers for editing the owner's build manual in its own
styles (Heading 1/2, Audio Guide Title/Body, Source Note, shaded tables,
callouts, centred figures with grey italic captions). Used by expand_v14.py.
"""
import copy, io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def open_doc(path):
    global doc, body
    doc = Document(path)
    body = doc.element.body
    return doc


def el_text(e):
    return ''.join(t.text or '' for t in e.iter(qn('w:t')))

def find_heading(text, style=None):
    for e in body.iterchildren():
        if e.tag == qn('w:p') and el_text(e).strip() == text:
            return e
    raise SystemExit('anchor not found: ' + text)

def before_break(h):
    """the page-break paragraph that opens a chapter, if there is one"""
    prev = h.getprevious()
    if prev is not None and prev.tag == qn('w:p') and prev.find('.//' + qn('w:br')) is not None and not el_text(prev).strip():
        return prev
    return h

def shd(el_pr, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    el_pr.append(s)

def borders(ppr, color):
    b = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        x = OxmlElement('w:' + side); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '6'); x.set(qn('w:color'), color)
        b.append(x)
    ppr.append(b)

class At:
    """inserts new blocks before an anchor element, in order"""
    def __init__(self, anchor): self.anchor = anchor
    def _put(self, el): self.anchor.addprevious(el); return el
    def _para(self, style=None):
        p = doc.add_paragraph(style=style)
        self._put(p._p)
        return p
    def pagebreak(self):
        p = self._para(); p.add_run().add_break(WD_BREAK.PAGE)
    def h1(self, t, newpage=True):
        if newpage: self.pagebreak()
        self._para('Heading 1').add_run(t)
    def h2(self, t): self._para('Heading 2').add_run(t)
    def p(self, t, italic=False, bold_lead=None):
        p = self._para()
        if bold_lead:
            r = p.add_run(bold_lead); r.bold = True
        r = p.add_run(t); r.italic = italic
        return p
    def bullets(self, items, mark='\u2022 '):
        for it in items:
            p = self._para()
            ppr = p._p.get_or_add_pPr()
            sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '40'); ppr.append(sp)
            ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '288'); ppr.append(ind)
            if isinstance(it, tuple):
                r = p.add_run(mark + it[0]); r.bold = True; r.font.size = Pt(10.5)
                r = p.add_run(it[1]); r.font.size = Pt(10.5)
            else:
                r = p.add_run(mark + it); r.font.size = Pt(10.5)
    def steps(self, items):
        for i, it in enumerate(items, 1):
            self.bullets([it if isinstance(it, tuple) else it], mark=f'{i}. ')
    def checks(self, items):
        for it in items:
            self._para().add_run('\u2610 ' + it)
    def audio(self, title, paras):
        p = self._para('Audio Guide Title'); ppr = p._p.get_or_add_pPr(); shd(ppr, '17365D'); borders(ppr, '17365D')
        p.add_run(title)
        for t in paras:
            p = self._para('Audio Guide Body'); ppr = p._p.get_or_add_pPr(); shd(ppr, 'EAF2F8'); borders(ppr, 'B4C6E7')
            p.add_run(t)
    def note(self, t, fill='FFF2CC', border='D6B656'):
        p = self._para('Source Note'); ppr = p._p.get_or_add_pPr(); shd(ppr, fill); borders(ppr, border)
        p.add_run(t)
    def source(self, t): self._para('Source Note').add_run(t)
    def _table(self, nrows, widths):
        total = 10224
        ws = [int(total * w / sum(widths)) for w in widths]
        t = doc.add_table(rows=nrows, cols=len(widths))
        tbl = t._tbl
        tblPr = tbl.tblPr
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
        for row in t.rows:
            for c, w in zip(row.cells, ws):
                c.width = w * 635  # dxa -> EMU
        self._put(tbl)
        return t, ws
    def table(self, header, rows, widths=None, fill='D9EAD3', size=9.5, bold_first=False):
        widths = widths or [1] * len(header)
        t, ws = self._table(len(rows) + 1, widths)
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; shd(c._tc.get_or_add_tcPr(), fill)
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                r = t.rows[i].cells[j].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
                if bold_first and j == 0: r.bold = True
        for row in t.rows:   # never split a row across pages
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        # keep the header with the first row; repeat it on a new page
        trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
        self._para()  # spacing after the table, as python-docx tables butt into the next block
        return t
    def callout(self, t, fill='FFF2CC', size=10.5):
        tb, _ = self._table(1, [1])
        c = tb.rows[0].cells[0]; shd(c._tc.get_or_add_tcPr(), fill)
        r = c.paragraphs[0].add_run(t); r.bold = True; r.font.size = Pt(size)
        self._para()
    def code(self, lines, size=8.5):
        tb, _ = self._table(1, [1])
        c = tb.rows[0].cells[0]; shd(c._tc.get_or_add_tcPr(), 'F2F2F2')
        p = c.paragraphs[0]
        for i, ln in enumerate(lines):
            r = p.add_run(ln); r.font.name = 'Consolas'; r.font.size = Pt(size)
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Consolas')
            if i < len(lines) - 1: r.add_break()
        self._para()
    def figure(self, path, caption, width=6.4):
        p = self._para(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(width))
        c = self._para(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(8)
        r = c.add_run(caption); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)


# ---- in-place edits ------------------------------------------------------------------
def paras_starting(prefix):
    return [e for e in body.iter(qn('w:p')) if el_text(e).startswith(prefix)]
def para(prefix):
    found = paras_starting(prefix)
    if len(found) != 1: raise SystemExit(f'expected one paragraph starting {prefix!r}, found {len(found)}')
    return found[0]
def set_text(p, text):
    """new text in the first run, keeping its formatting; other runs emptied"""
    ts = list(p.iter(qn('w:t')))
    ts[0].text = text
    ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    for t in ts[1:]: t.text = ''
def replace_in(p, old, new):
    full = el_text(p)
    if old not in full: raise SystemExit(f'{old!r} not in {full[:80]!r}')
    set_text(p, full.replace(old, new))
def table_with(text):
    found = [t for t in body.iter(qn('w:tbl')) if text in el_text(t)]
    if len(found) != 1: raise SystemExit(f'expected one table with {text!r}, found {len(found)}')
    return found[0]
def rows(tbl): return tbl.findall(qn('w:tr'))
def cells(tr): return tr.findall(qn('w:tc'))
def row_with(tbl, text):
    for tr in rows(tbl):
        if text in el_text(tr): return tr
    raise SystemExit(f'no row with {text!r}')
def set_cell(tc, text):
    ps = tc.findall(qn('w:p'))
    if list(ps[0].iter(qn('w:t'))): set_text(ps[0], text)
    else:
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = text; r.append(t); ps[0].append(r)
    for extra in ps[1:]: tc.remove(extra)
def add_row(tbl, values, after=None, like=None):
    src = like if like is not None else rows(tbl)[-1]
    new = copy.deepcopy(src)
    for tc, v in zip(cells(new), values): set_cell(tc, v)
    (after if after is not None else rows(tbl)[-1]).addnext(new)
    return new
def replace_figure(caption_prefix, png_path, caption=None, max_h_in=None):
    """swap the picture above a caption for a new PNG (same width, new height;
    with max_h_in, a tall picture is scaled down to that height instead)"""
    from PIL import Image
    cap = para(caption_prefix)
    pic = cap.getprevious()
    blip = next(pic.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))
    rid = blip.get(qn('r:embed'))
    part = doc.part.related_parts[rid]
    part._blob = open(png_path, 'rb').read()
    w, h = Image.open(png_path).size
    for ext in list(pic.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')) + \
               list(pic.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext')):
        cx = int(ext.get('cx')); cy = int(cx * h / w)
        if max_h_in and cy > max_h_in * 914400:
            cy = int(max_h_in * 914400); ext.set('cx', str(int(cy * w / h)))
        ext.set('cy', str(cy))
    if caption: set_text(cap, caption)
def table_hdr(prefix):
    """the table whose first row's text starts with prefix"""
    found = [t for t in body.iter(qn('w:tbl')) if el_text(rows(t)[0]).startswith(prefix)]
    if len(found) != 1: raise SystemExit(f'expected one table headed {prefix!r}, found {len(found)}')
    return found[0]
